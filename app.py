from flask import Flask, render_template, request, redirect, url_for, session, flash
import psycopg2
import psycopg2.extras
import resend
import cloudinary
import cloudinary.uploader
import africastalking
from itsdangerous import URLSafeTimedSerializer
import bcrypt
import os
from datetime import timedelta
from functools import wraps

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'ampath_secret_change_this_in_production')
app.permanent_session_lifetime = timedelta(minutes=30)

# ── PostgreSQL Configuration ──────────────────────────────────────────────────

DATABASE_URL = os.environ.get('DATABASE_URL')

def get_db():
    conn = psycopg2.connect(DATABASE_URL, cursor_factory=psycopg2.extras.RealDictCursor)
    return conn

# ── External Services ─────────────────────────────────────────────────────────

resend.api_key = os.environ.get("RESEND_API_KEY")

cloudinary.config(
    cloud_name=os.environ.get('CLOUDINARY_CLOUD_NAME'),
    api_key=os.environ.get('CLOUDINARY_API_KEY'),
    api_secret=os.environ.get('CLOUDINARY_API_SECRET')
)

africastalking.initialize(
    os.environ.get('AT_USERNAME', 'sandbox'),
    os.environ.get('AT_API_KEY')
)
sms = africastalking.SMS
s = URLSafeTimedSerializer(app.secret_key)

# ── Helper Functions ──────────────────────────────────────────────────────────

def auto_assign(incident_type):
    routing = {
        'Malware / Virus Detected': 'Victor',
        'Ransomware': 'Victor',
        'Phishing / Suspicious Email': 'Donald',
        'Unauthorised Access Attempt': 'Donald',
        'Data Breach / Leak': 'Donald',
        'Password Compromise': 'Donald',
        'Suspicious Network Activity': 'Alvin',
        'System Compromise': 'Denzel',
        'Lost / Stolen Device': 'Babu',
        'Other': 'Alvin',
    }
    return routing.get(incident_type, 'Alvin')

TEAM_EMAILS = {
    'Victor': 'njorogedunnant@gmail.com',
    'Donald': 'njorogedunnant@gmail.com',
    'Denzel': 'njorogedunnant@gmail.com',
    'Babu': 'njorogedunnant@gmail.com',
    'Alvin': 'njorogedunnant@gmail.com',
}

def send_assignment_email(assigned_to, incident_type, priority, incident_id):
    email = TEAM_EMAILS.get(assigned_to)
    if email:
        try:
            body = f"Hello {assigned_to},\n\nA new {priority} incident has been assigned to you.\n\nIncident #{incident_id}: {incident_type}\n\nPlease log in to view and resolve it:\nhttps://ampath-incident-portal.onrender.com\n\nAMPATH ICT Portal"
            resend.Emails.send({
                "from": "onboarding@resend.dev",
                "to": email,
                "subject": f"[{priority}] New Incident Assigned – {incident_type}",
                "text": body
            })
        except Exception as e:
            print(f"Email error: {e}")

def send_sms_alert(priority, incident_type, incident_id):
    if priority in ['P1', 'P2']:
        phone = os.environ.get('AT_PHONE', '')
        if phone:
            message = f"AMPATH ALERT: {priority} Incident #{incident_id} - {incident_type} has been reported. Immediate attention required!"
            try:
                sms.send(message, [phone])
            except Exception as e:
                print(f"SMS error: {e}")

def calculate_priority(urgency, impact):
    matrix = {
        ('High', 'High'): 'P1',
        ('High', 'Medium'): 'P2',
        ('Medium', 'High'): 'P2',
        ('High', 'Low'): 'P3',
        ('Medium', 'Medium'): 'P3',
        ('Low', 'High'): 'P3',
        ('Medium', 'Low'): 'P4',
        ('Low', 'Medium'): 'P4',
        ('Low', 'Low'): 'P5',
    }
    return matrix.get((urgency, impact), 'P5')

# ── Auth Decorators ───────────────────────────────────────────────────────────

def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            flash('Please log in to continue.', 'warning')
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated

def admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        if session.get('role') != 'admin':
            flash('Access denied. Admins only.', 'danger')
            return redirect(url_for('dashboard'))
        return f(*args, **kwargs)
    return decorated

def technician_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        if session.get('role') not in ['admin', 'technician']:
            flash('Access denied.', 'danger')
            return redirect(url_for('dashboard'))
        return f(*args, **kwargs)
    return decorated

# ── Routes ────────────────────────────────────────────────────────────────────

@app.route('/')
def index():
    if 'user_id' in session:
        return redirect(url_for('dashboard'))
    return redirect(url_for('login'))

# ── Auth ──────────────────────────────────────────────────────────────────────

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        email    = request.form.get('email', '').strip().lower()
        password = request.form.get('password', '')
        conn = get_db()
        cur = conn.cursor()
        cur.execute("SELECT * FROM users WHERE email = %s", (email,))
        user = cur.fetchone()
        cur.close(); conn.close()
        if user and bcrypt.checkpw(password.encode('utf-8'), user['password'].encode('utf-8')):
            session.permanent = True
            session['user_id'] = user['id']
            session['name']    = user['name']
            session['role']    = user['role']
            flash(f"Welcome back, {user['name']}!", 'success')
            return redirect(url_for('dashboard'))
        else:
            flash('Invalid email or password.', 'danger')
    return render_template('login.html')


@app.route('/logout')
def logout():
    session.clear()
    flash('You have been logged out.', 'info')
    return redirect(url_for('login'))

# ── Dashboard ─────────────────────────────────────────────────────────────────

@app.route('/dashboard')
@login_required
def dashboard():
    if session.get('role') == 'technician':
        return redirect(url_for('technician_dashboard'))
    conn = get_db()
    cur = conn.cursor()
    if session['role'] == 'admin':
        severity = request.args.get('severity', '')
        status   = request.args.get('status', '')
        query    = "SELECT i.*, u.name AS reporter_name FROM incidents i JOIN users u ON i.user_id = u.id WHERE 1=1"
        params   = []
        if severity:
            query += " AND i.severity = %s"
            params.append(severity)
        if status:
            query += " AND i.status = %s"
            params.append(status)
        query += " ORDER BY i.created_at DESC"
        cur.execute(query, params)
    else:
        cur.execute(
            "SELECT * FROM incidents WHERE user_id = %s ORDER BY created_at DESC",
            (session['user_id'],)
        )
    incidents = cur.fetchall()
    stats = {}
    if session['role'] == 'admin':
        cur.execute("SELECT COUNT(*) AS total FROM incidents")
        stats['total'] = cur.fetchone()['total']
        cur.execute("SELECT COUNT(*) AS open FROM incidents WHERE status = 'Open'")
        stats['open'] = cur.fetchone()['open']
        cur.execute("SELECT COUNT(*) AS critical FROM incidents WHERE severity = 'Critical'")
        stats['critical'] = cur.fetchone()['critical']
        cur.execute("SELECT COUNT(*) AS resolved FROM incidents WHERE status = 'Resolved'")
        stats['resolved'] = cur.fetchone()['resolved']
    cur.close(); conn.close()
    return render_template('dashboard.html', incidents=incidents, stats=stats)

# ── Report Incident ───────────────────────────────────────────────────────────

@app.route('/report', methods=['GET', 'POST'])
@login_required
def report_incident():
    if request.method == 'POST':
        incident_type = request.form.get('incident_type', '').strip()
        description   = request.form.get('description', '').strip()
        severity      = request.form.get('severity', '')
        location      = request.form.get('location', '').strip()
        urgency       = request.form.get('urgency', 'Low')
        impact        = request.form.get('impact', 'Low')
        priority      = calculate_priority(urgency, impact)
        assigned_to   = auto_assign(incident_type)
        photo_url     = None
        if 'photo' in request.files:
            photo = request.files['photo']
            if photo and photo.filename != '':
                allowed_extensions = {'jpg', 'jpeg', 'png', 'gif', 'webp'}
                file_ext = photo.filename.rsplit('.', 1)[-1].lower() if '.' in photo.filename else ''
                allowed_mimetypes = {'image/jpeg', 'image/png', 'image/gif', 'image/webp'}
                photo.seek(0, 2)
                file_size = photo.tell()
                photo.seek(0)
                if file_ext not in allowed_extensions:
                    flash('Invalid file type. Only JPG, PNG, GIF and WEBP images are allowed.', 'danger')
                    return render_template('report.html')
                elif file_size > 5 * 1024 * 1024:
                    flash('File too large. Maximum size is 5MB.', 'danger')
                    return render_template('report.html')
                elif photo.mimetype not in allowed_mimetypes:
                    flash('Invalid file type detected.', 'danger')
                    return render_template('report.html')
                else:
                    try:
                        upload_result = cloudinary.uploader.upload(
                            photo,
                            resource_type='image',
                            allowed_formats=['jpg', 'jpeg', 'png', 'gif', 'webp'],
                            max_bytes=5 * 1024 * 1024
                        )
                        photo_url = upload_result.get('secure_url')
                    except Exception as e:
                        print(f"Photo upload error: {e}")
                        flash('Photo upload failed. Please try again.', 'danger')
                        return render_template('report.html')
        if not all([incident_type, description, severity]):
            flash('Please fill in all required fields.', 'danger')
            return render_template('report.html')
        conn = get_db()
        cur = conn.cursor()
        cur.execute(
            "INSERT INTO incidents (user_id, incident_type, description, severity, location, status, urgency, impact, priority, assigned_to, photo) VALUES (%s, %s, %s, %s, %s, 'Open', %s, %s, %s, %s, %s) RETURNING id",
            (session['user_id'], incident_type, description, severity, location, urgency, impact, priority, assigned_to, photo_url)
        )
        incident_id = cur.fetchone()['id']
        conn.commit()
        cur.close(); conn.close()
        send_sms_alert(priority, incident_type, incident_id)
        send_assignment_email(assigned_to, incident_type, priority, incident_id)
        flash('Incident reported successfully! The ICT team has been notified.', 'success')
        return redirect(url_for('dashboard'))
    return render_template('report.html')

@app.route('/incident/<int:id>/update', methods=['POST'])
@admin_required
def update_incident(id):
    status   = request.form.get('status', '')
    assignee = request.form.get('assignee', '').strip()
    notes    = request.form.get('admin_notes', '').strip()
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "UPDATE incidents SET status = %s, assigned_to = %s, admin_notes = %s WHERE id = %s",
        (status, assignee, notes, id)
    )
    conn.commit()
    cur.close(); conn.close()
    flash('Incident updated successfully.', 'success')
    return redirect(url_for('dashboard'))


@app.route('/incident/<int:id>')
@login_required
def view_incident(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "SELECT i.*, u.name AS reporter_name, u.department FROM incidents i JOIN users u ON i.user_id = u.id WHERE i.id = %s",
        (id,)
    )
    incident = cur.fetchone()
    cur.close(); conn.close()
    if not incident:
        flash('Incident not found.', 'danger')
        return redirect(url_for('dashboard'))
    if session['role'] != 'admin' and incident['user_id'] != session['user_id']:
        flash('Access denied.', 'danger')
        return redirect(url_for('dashboard'))
    return render_template('view_incident.html', incident=incident)

# ── Forgot Password ───────────────────────────────────────────────────────────

@app.route('/forgot-password', methods=['GET', 'POST'])
def forgot_password():
    if request.method == 'POST':
        email = request.form.get('email', '').strip().lower()
        try:
            conn = get_db()
            cur = conn.cursor()
            cur.execute("SELECT * FROM users WHERE email = %s", (email,))
            user = cur.fetchone()
            cur.close(); conn.close()
            if user:
                token = s.dumps(email, salt='password-reset')
                reset_url = url_for('reset_password_page', token=token, _external=True)
                body = f"Hello {user['name']},\n\nClick the link below to reset your password (valid for 30 minutes):\n{reset_url}\n\nAMPATH ICT Team"
                resend.Emails.send({
                    "from": "onboarding@resend.dev",
                    "to": email,
                    "subject": "Reset Your AMPATH Portal Password",
                    "text": body
                })
        except Exception as e:
            print(f"Error: {e}")
            flash(f'Error: {str(e)}', 'danger')
            return render_template('forgot_password.html')
        flash('If that email exists, a reset link has been sent.', 'info')
        return redirect(url_for('login'))
    return render_template('forgot_password.html')

@app.route('/reset-password/<token>', methods=['GET', 'POST'])
def reset_password_page(token):
    try:
        email = s.loads(token, salt='password-reset', max_age=1800)
    except:
        flash('The reset link is invalid or has expired.', 'danger')
        return redirect(url_for('forgot_password'))
    if request.method == 'POST':
        password = request.form.get('password', '')
        if len(password) < 8:
            flash('Password must be at least 8 characters.', 'danger')
            return render_template('reset_password.html', token=token)
        hashed = bcrypt.hashpw(password.encode('utf-8'), bcrypt.gensalt())
        conn = get_db()
        cur = conn.cursor()
        cur.execute("UPDATE users SET password = %s WHERE email = %s", (hashed.decode('utf-8'), email))
        conn.commit()
        cur.close(); conn.close()
        flash('Password reset successfully! Please log in.', 'success')
        return redirect(url_for('login'))
    return render_template('reset_password.html', token=token)

# ── Trends ────────────────────────────────────────────────────────────────────

@app.route('/trends')
@admin_required
def trends():
    conn = get_db()
    cur = conn.cursor()

    # Incidents per month (PostgreSQL syntax)
    cur.execute("""
        SELECT TO_CHAR(MIN(created_at), 'Mon YYYY') as month,
               COUNT(*) as count
        FROM incidents
        GROUP BY TO_CHAR(created_at, 'YYYY-MM')
        ORDER BY TO_CHAR(created_at, 'YYYY-MM') DESC
        LIMIT 6
    """)
    monthly = cur.fetchall()

    cur.execute("""
        SELECT incident_type, COUNT(*) as count
        FROM incidents
        GROUP BY incident_type
        ORDER BY count DESC
        LIMIT 8
    """)
    by_type = cur.fetchall()

    cur.execute("SELECT severity, COUNT(*) as count FROM incidents GROUP BY severity")
    by_severity = cur.fetchall()

    cur.execute("SELECT status, COUNT(*) as count FROM incidents GROUP BY status")
    by_status = cur.fetchall()

    cur.execute("SELECT COUNT(*) as total FROM incidents")
    total = cur.fetchone()['total']

    cur.execute("SELECT COUNT(*) as resolved FROM incidents WHERE status = 'Resolved'")
    resolved = cur.fetchone()['resolved']

    cur.execute("""
        SELECT i.*, u.name AS reporter_name
        FROM incidents i
        JOIN users u ON i.user_id = u.id
        ORDER BY i.created_at DESC LIMIT 5
    """)
    recent = cur.fetchall()

    cur.close(); conn.close()
    return render_template('trends.html',
        monthly=monthly, by_type=by_type, by_severity=by_severity,
        by_status=by_status, total=total, resolved=resolved, recent=recent)

# ── Feedback ──────────────────────────────────────────────────────────────────

@app.route('/incident/<int:id>/feedback', methods=['GET', 'POST'])
@login_required
def feedback(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM incidents WHERE id = %s AND user_id = %s", (id, session['user_id']))
    incident = cur.fetchone()
    if not incident:
        flash('Incident not found or access denied.', 'danger')
        cur.close(); conn.close()
        return redirect(url_for('dashboard'))
    if incident['status'] != 'Resolved':
        flash('You can only give feedback on resolved incidents.', 'warning')
        cur.close(); conn.close()
        return redirect(url_for('dashboard'))
    if incident['feedback_rating']:
        flash('You have already submitted feedback for this incident.', 'info')
        cur.close(); conn.close()
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        rating  = request.form.get('rating')
        comment = request.form.get('comment', '').strip()
        cur.execute(
            "UPDATE incidents SET feedback_rating = %s, feedback_comment = %s WHERE id = %s",
            (rating, comment, id)
        )
        conn.commit()
        cur.close(); conn.close()
        flash('Thank you for your feedback!', 'success')
        return redirect(url_for('dashboard'))
    cur.close(); conn.close()
    return render_template('feedback.html', incident=incident)

# ── Knowledge Base ────────────────────────────────────────────────────────────

@app.route('/knowledge-base')
@login_required
def knowledge_base():
    search = request.args.get('search', '')
    incident_type = request.args.get('incident_type', '')
    conn = get_db()
    cur = conn.cursor()
    query = "SELECT * FROM knowledge_base WHERE 1=1"
    params = []
    if search:
        query += " AND (title ILIKE %s OR problem ILIKE %s OR solution ILIKE %s)"
        params.extend([f'%{search}%', f'%{search}%', f'%{search}%'])
    if incident_type:
        query += " AND incident_type = %s"
        params.append(incident_type)
    query += " ORDER BY created_at DESC"
    cur.execute(query, params)
    articles = cur.fetchall()
    cur.close(); conn.close()
    return render_template('knowledge_base.html', articles=articles, search=search, incident_type=incident_type)

@app.route('/knowledge-base/create', methods=['GET', 'POST'])
@admin_required
def create_kb_article():
    incident_id = request.args.get('incident_id')
    incident = None
    if incident_id:
        conn = get_db()
        cur = conn.cursor()
        cur.execute("SELECT * FROM incidents WHERE id = %s", (incident_id,))
        incident = cur.fetchone()
        cur.close(); conn.close()
    if request.method == 'POST':
        title         = request.form.get('title', '').strip()
        incident_type = request.form.get('incident_type', '').strip()
        problem       = request.form.get('problem', '').strip()
        solution      = request.form.get('solution', '').strip()
        inc_id        = request.form.get('incident_id') or None
        if not all([title, incident_type, problem, solution]):
            flash('All fields are required.', 'danger')
            return render_template('create_kb_article.html', incident=incident)
        conn = get_db()
        cur = conn.cursor()
        cur.execute(
            "INSERT INTO knowledge_base (title, incident_type, problem, solution, created_by, incident_id) VALUES (%s, %s, %s, %s, %s, %s)",
            (title, incident_type, problem, solution, session['name'], inc_id)
        )
        conn.commit()
        cur.close(); conn.close()
        flash('Knowledge Base article created successfully!', 'success')
        return redirect(url_for('knowledge_base'))
    return render_template('create_kb_article.html', incident=incident)

@app.route('/knowledge-base/<int:id>')
@login_required
def view_kb_article(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM knowledge_base WHERE id = %s", (id,))
    article = cur.fetchone()
    cur.close(); conn.close()
    if not article:
        flash('Article not found.', 'danger')
        return redirect(url_for('knowledge_base'))
    return render_template('view_kb_article.html', article=article)

@app.route('/knowledge-base/<int:id>/delete', methods=['POST'])
@admin_required
def delete_kb_article(id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("DELETE FROM knowledge_base WHERE id = %s", (id,))
    conn.commit()
    cur.close(); conn.close()
    flash('Article deleted.', 'success')
    return redirect(url_for('knowledge_base'))

# ── Technician Dashboard ──────────────────────────────────────────────────────

@app.route('/technician')
@technician_required
def technician_dashboard():
    conn = get_db()
    cur = conn.cursor()
    name   = session.get('name')
    status = request.args.get('status', '')
    query  = "SELECT i.*, u.name AS reporter_name FROM incidents i JOIN users u ON i.user_id = u.id WHERE i.assigned_to = %s"
    params = [name]
    if status:
        query += " AND i.status = %s"
        params.append(status)
    query += " ORDER BY i.created_at DESC"
    cur.execute(query, params)
    incidents = cur.fetchall()
    cur.execute("SELECT COUNT(*) AS total FROM incidents WHERE assigned_to = %s", (name,))
    total = cur.fetchone()['total']
    cur.execute("SELECT COUNT(*) AS open FROM incidents WHERE assigned_to = %s AND status = 'Open'", (name,))
    open_count = cur.fetchone()['open']
    cur.execute("SELECT COUNT(*) AS inprogress FROM incidents WHERE assigned_to = %s AND status = 'In Progress'", (name,))
    inprogress = cur.fetchone()['inprogress']
    cur.execute("SELECT COUNT(*) AS resolved FROM incidents WHERE assigned_to = %s AND status = 'Resolved'", (name,))
    resolved = cur.fetchone()['resolved']
    cur.close(); conn.close()
    return render_template('technician_dashboard.html', incidents=incidents,
        total=total, open_count=open_count, inprogress=inprogress, resolved=resolved)

@app.route('/technician/incident/<int:id>/update', methods=['POST'])
@technician_required
def technician_update_incident(id):
    status = request.form.get('status', '')
    notes  = request.form.get('admin_notes', '').strip()
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "UPDATE incidents SET status = %s, admin_notes = %s WHERE id = %s AND assigned_to = %s",
        (status, notes, id, session.get('name'))
    )
    conn.commit()
    cur.close(); conn.close()
    flash('Incident updated successfully.', 'success')
    return redirect(url_for('technician_dashboard'))

# ── User Management ───────────────────────────────────────────────────────────

@app.route('/users')
@admin_required
def manage_users():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT id, name, email, department, role, created_at FROM users ORDER BY role, name")
    users = cur.fetchall()
    cur.close(); conn.close()
    return render_template('manage_users.html', users=users)

@app.route('/users/create', methods=['GET', 'POST'])
@admin_required
def create_user():
    if request.method == 'POST':
        name     = request.form.get('name', '').strip()
        email    = request.form.get('email', '').strip().lower()
        password = request.form.get('password', '')
        dept     = request.form.get('department', '').strip()
        role     = request.form.get('role', 'staff')
        if not all([name, email, password, dept]):
            flash('All fields are required.', 'danger')
            return render_template('create_user.html')
        conn = get_db()
        cur = conn.cursor()
        cur.execute("SELECT id FROM users WHERE email = %s", (email,))
        if cur.fetchone():
            flash('An account with that email already exists.', 'danger')
            cur.close(); conn.close()
            return render_template('create_user.html')
        hashed = bcrypt.hashpw(password.encode('utf-8'), bcrypt.gensalt())
        cur.execute(
            "INSERT INTO users (name, email, password, department, role) VALUES (%s, %s, %s, %s, %s)",
            (name, email, hashed.decode('utf-8'), dept, role)
        )
        conn.commit()
        cur.close(); conn.close()
        flash(f'Account for {name} created successfully!', 'success')
        return redirect(url_for('manage_users'))
    return render_template('create_user.html')

@app.route('/users/<int:id>/delete', methods=['POST'])
@admin_required
def delete_user(id):
    if id == session['user_id']:
        flash('You cannot delete your own account.', 'danger')
        return redirect(url_for('manage_users'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute("DELETE FROM users WHERE id = %s", (id,))
    conn.commit()
    cur.close(); conn.close()
    flash('User deleted successfully.', 'success')
    return redirect(url_for('manage_users'))

@app.route('/users/<int:id>/role', methods=['POST'])
@admin_required
def change_role(id):
    role = request.form.get('role', 'staff')
    if id == session['user_id']:
        flash('You cannot change your own role.', 'danger')
        return redirect(url_for('manage_users'))
    conn = get_db()
    cur = conn.cursor()
    cur.execute("UPDATE users SET role = %s WHERE id = %s", (role, id))
    conn.commit()
    cur.close(); conn.close()
    flash('User role updated successfully.', 'success')
    return redirect(url_for('manage_users'))

# ── Export ────────────────────────────────────────────────────────────────────

@app.route('/export/excel')
@admin_required
def export_excel():
    import openpyxl
    from flask import make_response
    import io
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT i.id, i.incident_type, u.name AS reporter, i.severity, i.priority, i.status, i.assigned_to, i.location, i.description, i.created_at FROM incidents i JOIN users u ON i.user_id = u.id ORDER BY i.created_at DESC")
    incidents = cur.fetchall()
    cur.close(); conn.close()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Incidents"
    headers = ["ID", "Type", "Reporter", "Severity", "Priority", "Status", "Assigned To", "Location", "Description", "Date"]
    from openpyxl.styles import Font, PatternFill, Alignment
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0F4C81")
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row, inc in enumerate(incidents, 2):
        values = [inc["id"], inc["incident_type"], inc["reporter"], inc["severity"], inc["priority"], inc["status"], inc["assigned_to"], inc["location"], inc["description"], str(inc["created_at"])]
        for col, value in enumerate(values, 1):
            ws.cell(row=row, column=col, value=value)
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    response = make_response(output.read())
    response.headers["Content-Disposition"] = "attachment; filename=AMPATH_Incidents.xlsx"
    response.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return response

@app.route("/export/pdf")
@admin_required
def export_pdf():
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from flask import make_response
    import io, datetime
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT i.id, i.incident_type, u.name AS reporter, i.severity, i.priority, i.status, i.assigned_to, i.created_at FROM incidents i JOIN users u ON i.user_id = u.id ORDER BY i.created_at DESC")
    incidents = cur.fetchall()
    cur.close(); conn.close()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), topMargin=1.5*cm, bottomMargin=1.5*cm, leftMargin=1.5*cm, rightMargin=1.5*cm)
    story = []
    story.append(Paragraph("AMPATH Incident Reports", ParagraphStyle("title", fontSize=18, textColor=colors.HexColor("#0F4C81"), alignment=TA_CENTER, fontName="Helvetica-Bold", spaceAfter=8)))
    story.append(Paragraph(f"Generated on {datetime.datetime.now().strftime('%d %B %Y %H:%M')}", ParagraphStyle("sub", fontSize=10, textColor=colors.HexColor("#6B7280"), alignment=TA_CENTER, fontName="Helvetica", spaceAfter=16)))
    data = [["#", "Type", "Reporter", "Severity", "Priority", "Status", "Assigned To", "Date"]]
    for inc in incidents:
        data.append([str(inc["id"]), inc["incident_type"] or "", inc["reporter"] or "", inc["severity"] or "", inc["priority"] or "P5", inc["status"] or "", inc["assigned_to"] or "", str(inc["created_at"])[:10]])
    t = Table(data, colWidths=[1*cm, 5*cm, 3.5*cm, 2.5*cm, 2*cm, 3*cm, 3.5*cm, 3*cm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#0F4C81")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTNAME", (0,1), (-1,-1), "Helvetica"),
        ("FONTSIZE", (0,0), (-1,-1), 8),
        ("ROWPADDING", (0,0), (-1,-1), 5),
        ("GRID", (0,0), (-1,-1), 0.5, colors.HexColor("#E5E7EB")),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#E8F0FB")]),
        ("ALIGN", (0,0), (-1,-1), "CENTER"),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
    ]))
    story.append(t)
    doc.build(story)
    buffer.seek(0)
    response = make_response(buffer.read())
    response.headers["Content-Disposition"] = "attachment; filename=AMPATH_Incidents.pdf"
    response.headers["Content-Type"] = "application/pdf"
    return response

if __name__ == '__main__':
    app.run(debug=True)
